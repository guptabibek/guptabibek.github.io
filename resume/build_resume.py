"""Builds the ATS-friendly résumé from one content source.

Outputs:
  resume/resume.html            -> printed to assets/Bibek-Gupta-Resume.pdf by print_pdf.mjs
  assets/Bibek-Gupta-Resume.docx

ATS rules this layout follows: single column, standard section names, plain text
(no tables, columns, text boxes, icons or header/footer content), and every date
printed inline on the same line as its role so parsers keep them together.

Usage: python3 resume/build_resume.py && node resume/print_pdf.mjs
"""
import re
from html import escape as _escape
from pathlib import Path


def escape(text: str) -> str:
    # Keep hyphenated words on one line so PDF text extraction never splits them ("sign-" / "on").
    return re.sub(r"(\S+-\S+)", r"<span class='nw'>\1</span>", _escape(text))


ROOT = Path(__file__).resolve().parent.parent

NAME = "Bibek Gupta"
HEADLINE = "Senior Software Engineer | Technical Lead | .NET, Azure & System Architecture"
CONTACT = [
    "Kathmandu, Nepal",
    "+977 9845910191",
    "guptabibek166@gmail.com",
    "linkedin.com/in/itsmebibek",
    "github.com/guptabibek",
    "guptabibek.github.io",
]

SUMMARY = (
    "Senior Software Engineer and Technical Lead with 8+ years building production SaaS, ERP, and "
    "enterprise systems in C#, .NET Core, SQL Server, and Microsoft Azure. Currently leading a team of "
    "7-8 engineers rewriting ERISApedia, an ERISA compliance-research platform now part of Ascensus, from legacy "
    "PHP to .NET Core, Angular, and MySQL on Azure and GCP, and shipping its AI Search (Ask ERISA) and Form 5500 "
    "data mining. Previously Senior Engineer and Team Lead on PensionPro, a SaaS platform used by 400+ "
    "third-party administrator firms managing 230,000+ U.S. retirement plans. Founder of SastoRent and "
    "builder of ForecastPro, an AI reporting SaaS in production with retail pharmacy chains in India."
)

SKILLS = [
    ("Languages", "C#, TypeScript, JavaScript, SQL, PHP"),
    ("Backend", ".NET 8, .NET Core, ASP.NET Core Web API, Entity Framework Core, Dapper, Node.js, NestJS, "
                "Express, REST APIs, WebSockets, SignalR, background jobs (Quartz.NET, BullMQ)"),
    ("Architecture", "System design, Clean Architecture, microservices, event-driven architecture, serverless, "
                     "multi-tenant SaaS, legacy modernization, API design, RBAC, high availability"),
    ("Cloud & DevOps", "Microsoft Azure (Functions, Logic Apps, Service Bus, Event Grid, App Service, Storage, "
                       "Azure AD), Google Cloud Platform, Docker, Azure DevOps, GitHub Actions, CI/CD, Linux, Nginx"),
    ("Databases", "SQL Server, PostgreSQL, MySQL, Redis, SQLite, Prisma, query and schema optimization"),
    ("Frontend", "Angular, React, Next.js, TypeScript, Tailwind CSS, PrimeNG, Ag-Grid"),
    ("Security & Integrations", "SAML 2.0 SSO, OAuth 2.0 / OpenID Connect, JWT, AES encryption, Stripe, DocuSign, "
                                "QuickBooks Online, Braintree, Zendesk, Google Sheets & Drive APIs"),
    ("AI & Data", "LLM integration (OpenAI-compatible APIs), LLM-powered search, natural language to SQL (NLQ), semantic layer design, "
                  "token metering, time-series forecasting"),
    ("Testing", "Unit and integration testing, Jest, Vitest, Playwright"),
    ("Leadership", "Team leadership, code review, design review, coding standards, mentoring, "
                   "requirements analysis, collaboration with U.S. stakeholders"),
]

# (title, company line, dates, intro, [(sub-heading, dates, bullets)] or bullets)
EXPERIENCE = [
    {
        "title": "Technical Lead / Senior Software Engineer",
        "company": "DolphinDive Technology (client: AmericanTCS, acquired by Ascensus in 2026) | Kathmandu, Nepal",
        "dates": "Jun 2022 - Present",
        "intro": "Engineering team for U.S. retirement-industry SaaS products (PensionPro, ERISApedia), now part of "
                 "Ascensus, which serves 16M+ savers and $930B+ in assets under administration.",
        "groups": [
            ("ERISApedia - Team Lead", "Nov 2025 - Present", [
                "Lead a team of 7-8 engineers on the ground-up rewrite of ERISApedia, an ERISA compliance-research "
                "platform, replacing a legacy PHP application with .NET Core, Angular, and MySQL.",
                "Team delivered AI Search (Ask ERISA), launched July 2026: LLM-powered answers across nine ERISA "
                "compliance book titles for advisors, auditors, and TPAs.",
                "Team built Form 5500 data mining: search across filings and attachments with hundreds of criteria, "
                "benchmarking, and map-based sales prospecting.",
                "Own the target architecture, technology selection, and migration path off the legacy system; "
                "designed the cloud architecture across Microsoft Azure and Google Cloud Platform.",
                "Made the new platform faster than the legacy PHP system by rewriting UI components, optimizing "
                "slow SQL queries, and adding API-layer caching.",
                "Partner with U.S. stakeholders and engineers on architecture, scope, and sequencing to keep the "
                "rewrite aligned with product and compliance requirements.",
            ]),
            ("PensionPro - Senior Software Engineer, then Team Lead (2 years)", "Jun 2022 - Nov 2025", [
                "Delivered features across the full stack (C#, .NET Core, Angular, SQL Server, Redis) of a pension "
                "administration SaaS used by 400+ TPA firms administering 230,000+ retirement plans.",
                "Moved long-running and integration-heavy work off the request path using Azure Service Bus, Storage "
                "queues and tables, Azure Functions (all trigger types), Logic Apps, and Event Grid, so background "
                "processing scales independently of user traffic.",
                "Integrated DocuSign, Zendesk, ftwilliam.com, QuickBooks Online, and Braintree; implemented SAML 2.0 "
                "single sign-on and identity management.",
                "Established the CI/CD and release strategy on Azure DevOps: build and release pipelines, environment "
                "promotion, and deployment gates, supporting production releases every 2 to 4 weeks.",
                "As Team Lead, set code review, design review, and coding standards and mentored up to 4 engineers "
                "plus QA at a time; served as "
                "primary engineering liaison translating pension-compliance requirements into technical designs.",
            ]),
        ],
    },
    {
        "title": "Senior Software Engineer (Freelance Consultant)",
        "company": "Bharuwa Solutions (India)",
        "dates": "Jun 2022 - Jun 2025",
        "intro": "Enterprise software products: ERP, fintech, distribution, and agritech platforms.",
        "bullets": [
            "Architected QueryGen, a domain-driven ASP.NET Core engine used in production that generates dynamic "
            "queries and reports with per-client database routing, eliminating hand-written SQL for routine reporting.",
            "Re-engineered ERP accounting, payroll, and invoicing modules for multi-tenant operation, letting one "
            "deployment serve multiple client organizations with isolated data.",
            "Automated GST reporting and filing against the government e-portal, replacing a manual compliance process.",
            "Built inventory and logistics tracking for a Warehouse Management System with multi-level category hierarchies.",
        ],
    },
    {
        "title": "Software Engineer",
        "company": "IMS Software Pvt. Ltd. | Kathmandu, Nepal",
        "dates": "Jun 2019 - May 2022",
        "intro": "Retail and distribution platforms used by 20,000+ businesses on 100,000+ POS terminals "
                 "across Nepal, India, and Southeast Asia.",
        "bullets": [
            "Engineered a custom ERP for Patanjali Ayurved (India), the company's flagship cross-border enterprise "
            "implementation on one of India's largest FMCG distribution networks.",
            "Built and owned production retail POS and distribution management systems: inventory, customers, "
            "sales reporting, and multi-level distribution hierarchies.",
            "Integrated accounting workflows into the POS and administered the production Linux servers behind it.",
        ],
    },
    {
        "title": "Software Engineer",
        "company": "Softweb Developers | Kathmandu, Nepal",
        "dates": "Apr 2018 - Mar 2019",
        "bullets": [
            "Designed the front-end architecture and a reusable component library shared across applications.",
            "Optimized relational schemas for query performance and added WebSocket-based real-time features.",
        ],
    },
]

PROJECTS = [
    ("ForecastPro - Multi-tenant forecasting & AI reporting SaaS", "Sole engineer, in production", [
        "In production with a few large retail pharmacy chains in India: syncs their ERP data, forecasts demand, "
        "and answers business questions asked in plain English.",
        "Built a natural-language reporting (NLQ) engine: an LLM maps questions to a catalog-constrained semantic "
        "query that is compiled to parameterized SQL with tenant and branch row-level security and SQL safety "
        "checks, so the model never writes SQL.",
        "Built an AI insights engine with 12 pluggable providers (stock-out risk, churn, dead stock, discount "
        "anomalies, business health score), Redis-cached dashboards, and scheduled generation.",
        "Built prepaid AI credit billing: append-only ledger enforced by database triggers, reservation-based "
        "charging, per-tenant pricing and spend limits, and signature-verified, idempotent Stripe webhooks.",
        "Built a resumable ERP data sync pipeline with queue-based, per-tenant worker concurrency.",
    ]),
    ("SastoRent - Founder & Lead Engineer", "sastorent.com", [
        "Building a marketplace where people rent out properties and everyday items, and where service providers "
        "such as plumbers, carpenters, and tutors register and sell their services.",
    ]),
    ("LearnSphere - Multi-tenant learning management system", "Lead engineer", [
        "SAML and OpenID Connect SSO, Stripe payments, queue-based exports, Redis caching, attendance and "
        "marksheets, and Prometheus metrics across 112 database tables; deployed with Docker.",
    ]),
    ("DriveNow and Futsal Booking - Booking platforms", "Freelance", [
        "Car rental and futsal court booking platforms with online payments, availability checks that prevent "
        "double bookings, role-based admin panels, and customer dashboards.",
    ]),
]

EDUCATION = [("BSc Computer Science", "Tribhuvan University", "2018")]


def build_html() -> str:
    out = []
    out.append(f"<header><h1>{escape(NAME)}</h1><p class='headline'>{escape(HEADLINE)}</p>"
               f"<p class='contact'>{' | '.join(escape(c) for c in CONTACT)}</p></header>")
    out.append(f"<h2>Summary</h2><p>{escape(SUMMARY)}</p>")
    out.append("<h2>Technical Skills</h2><ul class='skills'>" + "".join(
        f"<li><b>{escape(k)}:</b> {escape(v)}</li>" for k, v in SKILLS) + "</ul>")
    out.append("<h2>Professional Experience</h2>")
    for job in EXPERIENCE:
        out.append(f"<h3>{escape(job['title'])} | {escape(job['dates'])}</h3>")
        out.append(f"<p class='org'>{escape(job['company'])}</p>")
        if job.get("intro"):
            out.append(f"<p class='intro'>{escape(job['intro'])}</p>")
        for sub, dates, bullets in job.get("groups", []):
            out.append(f"<h4>{escape(sub)} | {escape(dates)}</h4>")
            out.append("<ul>" + "".join(f"<li>{escape(b)}</li>" for b in bullets) + "</ul>")
        if job.get("bullets"):
            out.append("<ul>" + "".join(f"<li>{escape(b)}</li>" for b in job["bullets"]) + "</ul>")
    out.append("<h2>Projects</h2>")
    for title, meta, bullets in PROJECTS:
        out.append(f"<h4>{escape(title)} | {escape(meta)}</h4>")
        out.append("<ul>" + "".join(f"<li>{escape(b)}</li>" for b in bullets) + "</ul>")
    out.append("<h2>Education</h2>")
    for degree, school, year in EDUCATION:
        out.append(f"<p><b>{escape(degree)}</b> | {escape(school)} | {escape(year)}</p>")

    css = """
    @page { size: A4; margin: 13mm 15mm; }
    body { margin: 0; font-family: 'Liberation Sans', Arial, Helvetica, sans-serif; font-size: 9.6pt; line-height: 1.32; color: #111; }
    h1 { margin: 0; font-size: 20pt; letter-spacing: .02em; }
    .headline { margin: 1pt 0 0; font-size: 10.5pt; font-weight: bold; color: #1f3a70; }
    .contact { margin: 2pt 0 0; font-size: 9pt; }
    h2 { margin: 8pt 0 3pt; padding-bottom: 1.5pt; border-bottom: 1px solid #1f3a70; color: #1f3a70; font-size: 10.5pt; text-transform: uppercase; letter-spacing: .06em; }
    h3 { margin: 6pt 0 0; font-size: 10pt; }
    h4 { margin: 4pt 0 1pt; font-size: 9.6pt; color: #1f3a70; }
    p { margin: 0; }
    .nw { white-space: nowrap; }
    .org { font-weight: bold; }
    .intro { font-style: italic; margin-top: 1pt; }
    ul { margin: 1pt 0 0; padding-left: 13pt; }
    li { margin: 0 0 1.2pt; }
    ul.skills { list-style: none; padding-left: 0; }
    h2, h3, h4 { break-after: avoid; }
    li { break-inside: avoid; }
    """
    return (f"<!doctype html><html lang='en'><head><meta charset='utf-8'><title>{_escape(NAME)} - Resume</title>"
            f"<style>{css}</style></head><body>{''.join(out)}</body></html>")


def build_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    blue = RGBColor(0x1F, 0x3A, 0x70)
    doc = Document()
    for s in doc.sections:
        s.page_height, s.page_width = Mm(297), Mm(210)
        s.top_margin = s.bottom_margin = Mm(13)
        s.left_margin = s.right_margin = Mm(15)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(1)
    normal.paragraph_format.line_spacing = 1.05

    def para(text="", bold=False, italic=False, size=None, color=None, after=None, before=None, style=None):
        p = doc.add_paragraph(style=style)
        if text:
            r = p.add_run(text)
            r.bold, r.italic = bold, italic
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
        if after is not None: p.paragraph_format.space_after = Pt(after)
        if before is not None: p.paragraph_format.space_before = Pt(before)
        return p

    def heading(text):
        # Real Heading 1 style so ATS and screen readers detect sections.
        h = doc.add_heading(text.upper(), level=1)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        for r in h.runs:
            r.font.size, r.font.color.rgb, r.font.name = Pt(11), blue, "Calibri"

    def bullets(items):
        for b in items:
            p = para(b, style="List Bullet", after=1)
            p.paragraph_format.left_indent = Mm(5)

    para(NAME, bold=True, size=20, after=0)
    para(HEADLINE, bold=True, size=11, color=blue, after=0)
    para(" | ".join(CONTACT), size=9.5, after=2)

    heading("Summary")
    para(SUMMARY)

    heading("Technical Skills")
    for k, v in SKILLS:
        p = para(after=1)
        p.add_run(f"{k}: ").bold = True
        p.add_run(v)

    heading("Professional Experience")
    for job in EXPERIENCE:
        para(f"{job['title']} | {job['dates']}", bold=True, size=10.5, before=5, after=0)
        para(job["company"], bold=True, after=0)
        if job.get("intro"):
            para(job["intro"], italic=True, after=1)
        for sub, dates, items in job.get("groups", []):
            para(f"{sub} | {dates}", bold=True, color=blue, before=3, after=0)
            bullets(items)
        if job.get("bullets"):
            bullets(job["bullets"])

    heading("Projects")
    for title, meta, items in PROJECTS:
        para(f"{title} | {meta}", bold=True, color=blue, before=3, after=0)
        bullets(items)

    heading("Education")
    for degree, school, year in EDUCATION:
        p = para()
        p.add_run(degree).bold = True
        p.add_run(f" | {school} | {year}")

    doc.core_properties.author = NAME
    doc.core_properties.title = f"{NAME} - Resume"
    doc.save(path)


if __name__ == "__main__":
    (ROOT / "resume" / "resume.html").write_text(build_html(), encoding="utf-8")
    build_docx(ROOT / "assets" / "Bibek-Gupta-Resume.docx")
    print("wrote resume/resume.html and assets/Bibek-Gupta-Resume.docx")
